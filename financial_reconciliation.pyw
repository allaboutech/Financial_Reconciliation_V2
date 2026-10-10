import csv
import os
import sqlite3
import threading
import queue
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

try:
    import openpyxl
except ImportError:
    openpyxl = None

APP_DIR = Path.home() / 'FinancialReconciliationToolkit'
APP_DIR.mkdir(exist_ok=True)
DB = APP_DIR / 'reconciliation.sqlite3'

SCHEMA = '''
CREATE TABLE IF NOT EXISTS runs (id INTEGER PRIMARY KEY AUTOINCREMENT, created_at TEXT, gl_file TEXT, sub_file TEXT, gl_count INTEGER, sub_count INTEGER, matched INTEGER, exceptions INTEGER, gl_total TEXT, sub_total TEXT);
CREATE TABLE IF NOT EXISTS transactions (id INTEGER PRIMARY KEY AUTOINCREMENT, run_id INTEGER, source TEXT, transaction_id TEXT, account TEXT, txn_date TEXT, amount_cents INTEGER);
CREATE TABLE IF NOT EXISTS exceptions (id INTEGER PRIMARY KEY AUTOINCREMENT, run_id INTEGER, transaction_id TEXT, account TEXT, issue TEXT, gl_count INTEGER, sub_count INTEGER, gl_cents INTEGER, sub_cents INTEGER, status TEXT DEFAULT 'Open', note TEXT DEFAULT '', reviewed_by TEXT DEFAULT '', reviewed_at TEXT DEFAULT '');
'''

def connection():
    c = sqlite3.connect(DB, timeout=15)
    c.row_factory = sqlite3.Row
    return c

with connection() as c:
    c.executescript(SCHEMA)

def money(cents):
    return f'{Decimal(cents) / 100:,.2f}'

def parse_amount(value):
    s = str(value if value is not None else '').strip().replace(',', '').replace('$', '')
    if s.startswith('(') and s.endswith(')'):
        s = '-' + s[1:-1]
    try:
        return int((Decimal(s) * 100).quantize(Decimal('1')))
    except (InvalidOperation, ValueError):
        raise ValueError(f'Invalid amount: {value!r}')

def normalize_date(value):
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d')
    s = str(value if value is not None else '').strip()
    for fmt in ('%Y-%m-%d', '%m/%d/%Y', '%m/%d/%y', '%Y/%m/%d', '%m-%d-%Y'):
        try:
            return datetime.strptime(s.split(' ')[0], fmt).strftime('%Y-%m-%d')
        except ValueError:
            pass
    return s

def load_rows(path):
    if path.lower().endswith('.csv'):
        with open(path, newline='', encoding='utf-8-sig') as f:
            raw = list(csv.reader(f))
    elif path.lower().endswith('.xlsx'):
        if openpyxl is None:
            raise RuntimeError('Install openpyxl: pip install openpyxl')
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            raw = list(wb.active.values)
        finally:
            wb.close()
    else:
        raise ValueError('Choose a CSV or XLSX file')
    if not raw:
        raise ValueError(f'Empty file: {path}')
    names = [str(x or '').strip().lower().replace(' ', '_') for x in raw[0]]
    aliases = {'transaction_id': ('transaction_id', 'transactionid', 'transaction', 'txn_id', 'id'), 'account': ('account', 'account_number', 'gl_account'), 'amount': ('amount', 'transaction_amount', 'debit_credit_amount'), 'date': ('date', 'transaction_date', 'posting_date', 'txn_date')}
    indices = {}
    for k, choices in aliases.items():
        indices[k] = next((names.index(n) for n in choices if n in names), None)
        if indices[k] is None:
            raise ValueError(f'{Path(path).name}: missing column {k}; headers: {names}')
    rows = []
    for line, values in enumerate(raw[1:], 2):
        if not values or all(v is None or str(v).strip() == '' for v in values):
            continue
        def field(k):
            idx = indices[k]
            return values[idx] if idx < len(values) else None
        tid = str(field('transaction_id') or '').strip()
        acc = str(field('account') or '').strip()
        if not tid or not acc:
            raise ValueError(f'{Path(path).name}, row {line}: transaction_id and account required')
        rows.append({'transaction_id': tid, 'account': acc, 'txn_date': normalize_date(field('date')), 'amount_cents': parse_amount(field('amount'))})
    return rows

def reconcile(gl, sub):
    from collections import defaultdict
    a, b = defaultdict(list), defaultdict(list)
    for r in gl:
        a[(r['transaction_id'], r['account'])].append(r)
    for r in sub:
        b[(r['transaction_id'], r['account'])].append(r)
    exceptions, matched = [], 0
    for key in sorted(set(a) | set(b)):
        left, right = a[key], b[key]
        ga = sum(x['amount_cents'] for x in left)
        sa = sum(x['amount_cents'] for x in right)
        if len(left) > 1 or len(right) > 1:
            issue = 'DUPLICATE_KEY'
        elif not left:
            issue = 'MISSING_GL'
        elif not right:
            issue = 'MISSING_SUBLEDGER'
        elif ga != sa:
            issue = 'AMOUNT_MISMATCH'
        elif left[0]['txn_date'] != right[0]['txn_date']:
            issue = 'DATE_MISMATCH'
        else:
            matched += 1
            continue
        exceptions.append((key[0], key[1], issue, len(left), len(right), ga, sa))
    return matched, exceptions

class Application(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title('Financial Reconciliation Automation | Controller Toolkit V2')
        self.geometry('1160x760')
        self.minsize(920, 620)
        self.gl = tk.StringVar()
        self.sub = tk.StringVar()
        self.run_id = tk.StringVar()
        self.status = tk.StringVar(value='Ready. Import GL and subledger files to begin.')
        self.q = queue.Queue()
        self.configure(padx=12, pady=10)
        ttk.Label(self, text='Financial Reconciliation Automation', font=('Segoe UI', 19, 'bold')).pack(anchor='w')
        ttk.Label(self, text='Controller Toolkit V2  |  Local SQLite  |  Audit-ready exception tracking').pack(anchor='w', pady=(0, 10))
        self.tabs = ttk.Notebook(self)
        self.tabs.pack(fill='both', expand=True)
        self.frames = {}
        for name in ('Dashboard', 'Reconciliation', 'SQL Analysis', 'Exception Management', 'Audit Reports'):
            frame = ttk.Frame(self.tabs, padding=12)
            self.tabs.add(frame, text=name)
            self.frames[name] = frame
        self.make_dashboard()
        self.make_reconciliation()
        self.make_sql()
        self.make_exceptions()
        self.make_reports()
        ttk.Label(self, textvariable=self.status).pack(anchor='w', pady=(7, 0))
        self.refresh_runs()
        self.after(100, self.poll)

    def table(self, parent, cols, height=14):
        frame = ttk.Frame(parent)
        frame.pack(fill='both', expand=True, pady=8)
        tree = ttk.Treeview(frame, columns=cols, show='headings', height=height)
        for col in cols:
            tree.heading(col, text=col.replace('_', ' ').title())
            tree.column(col, width=120, stretch=True)
        y = ttk.Scrollbar(frame, orient='vertical', command=tree.yview)
        x = ttk.Scrollbar(frame, orient='horizontal', command=tree.xview)
        tree.configure(yscrollcommand=y.set, xscrollcommand=x.set)
        tree.grid(row=0, column=0, sticky='nsew')
        y.grid(row=0, column=1, sticky='ns')
        x.grid(row=1, column=0, sticky='ew')
        frame.rowconfigure(0, weight=1)
        frame.columnconfigure(0, weight=1)
        return tree

    def make_dashboard(self):
        f = self.frames['Dashboard']
        ttk.Label(f, text='Reconciliation history', font=('Segoe UI', 14, 'bold')).pack(anchor='w')
        ttk.Label(f, text='Balances reflect uploaded signed transaction amounts, not independently validated account balances.').pack(anchor='w')
        ttk.Button(f, text='Refresh Dashboard', command=self.refresh_runs).pack(anchor='w', pady=8)
        self.runs_tree = self.table(f, ('run_id', 'created_at', 'gl_rows', 'sub_rows', 'matched_keys', 'exceptions', 'gl_total', 'sub_total', 'difference'))
        self.runs_tree.bind('<<TreeviewSelect>>', self.choose_run)

    def make_reconciliation(self):
        f = self.frames['Reconciliation']
        for label, variable in [('General Ledger', self.gl), ('Subledger', self.sub)]:
            r = ttk.Frame(f)
            r.pack(fill='x', pady=5)
            ttk.Label(r, text=label, width=18).pack(side='left')
            ttk.Entry(r, textvariable=variable).pack(side='left', fill='x', expand=True, padx=5)
            ttk.Button(r, text='Browse Excel / CSV', command=lambda v=variable: self.browse(v)).pack(side='left')
        ttk.Button(f, text='Run Reconciliation', command=self.start_reconcile).pack(anchor='w', pady=8)
        self.result_label = ttk.Label(f, text='No reconciliation run selected')
        self.result_label.pack(anchor='w')
        self.recon_tree = self.table(f, ('transaction_id', 'account', 'issue', 'gl_count', 'sub_count', 'gl_amount', 'sub_amount'))

    def make_sql(self):
        f = self.frames['SQL Analysis']
        ttk.Label(f, text='Read-only SQLite analysis (SELECT and WITH only)', font=('Segoe UI', 13, 'bold')).pack(anchor='w')
        ttk.Label(f, text='Tables: runs, transactions, exceptions. Filter with run_id for a specific reconciliation.').pack(anchor='w')
        self.sql_text = tk.Text(f, height=6, wrap='none', font=('Consolas', 10))
        self.sql_text.pack(fill='x', pady=8)
        self.sql_text.insert('1.0', 'SELECT run_id, issue, COUNT(*) AS exception_count\nFROM exceptions\nGROUP BY run_id, issue\nORDER BY run_id DESC, exception_count DESC;')
        ttk.Button(f, text='Run Read-only Query', command=self.run_sql).pack(anchor='w')
        self.sql_result = self.table(f, ('result',), 13)

    def make_exceptions(self):
        f = self.frames['Exception Management']
        bar = ttk.Frame(f)
        bar.pack(fill='x')
        ttk.Label(bar, text='Run ID:').pack(side='left')
        ttk.Entry(bar, textvariable=self.run_id, width=10).pack(side='left', padx=5)
        ttk.Button(bar, text='Load Exceptions', command=self.load_exceptions).pack(side='left')
        self.exc_tree = self.table(f, ('id', 'transaction_id', 'account', 'issue', 'difference', 'status', 'reviewed_by', 'reviewed_at'), 11)
        bottom = ttk.Frame(f)
        bottom.pack(fill='x')
        ttk.Label(bottom, text='Status').grid(row=0, column=0, sticky='w')
        self.exc_status = ttk.Combobox(bottom, values=['Open', 'Investigating', 'Resolved', 'Accepted Difference'], state='readonly', width=20)
        self.exc_status.set('Investigating')
        self.exc_status.grid(row=0, column=1, padx=5)
        ttk.Label(bottom, text='Reviewer').grid(row=0, column=2)
        self.reviewer = ttk.Entry(bottom, width=22)
        self.reviewer.grid(row=0, column=3, padx=5)
        ttk.Label(bottom, text='Explanation / evidence reference').grid(row=1, column=0, sticky='w', pady=7)
        self.note = ttk.Entry(bottom, width=75)
        self.note.grid(row=1, column=1, columnspan=3, sticky='ew', pady=7)
        ttk.Button(bottom, text='Save Selected Exception Review', command=self.save_review).grid(row=2, column=0, columnspan=4, sticky='w')
        bottom.columnconfigure(3, weight=1)

    def make_reports(self):
        f = self.frames['Audit Reports']
        ttk.Label(f, text='Audit-ready Excel workbook', font=('Segoe UI', 14, 'bold')).pack(anchor='w')
        ttk.Label(f, text='Exports run summary, original imported transactions, and exception review records.').pack(anchor='w', pady=6)
        ttk.Label(f, text='Select a run on the Dashboard or enter its ID on the Exception Management tab.').pack(anchor='w')
        ttk.Button(f, text='Export Selected Run to Excel', command=self.export_report).pack(anchor='w', pady=12)
        ttk.Label(f, text='Review statuses and notes are management documentation; the application does not post journal entries.').pack(anchor='w')

    def browse(self, variable):
        p = filedialog.askopenfilename(filetypes=[('Excel / CSV', '*.xlsx *.csv'), ('All files', '*.*')])
        if p:
            variable.set(p)

    def start_reconcile(self):
        gl, sub = self.gl.get(), self.sub.get()
        if not os.path.isfile(gl) or not os.path.isfile(sub):
            messagebox.showerror('Missing files', 'Select existing GL and subledger files.')
            return
        self.status.set('Processing reconciliation...')
        threading.Thread(target=self.worker, args=(gl, sub), daemon=True).start()

    def worker(self, glpath, subpath):
        try:
            gl, sub = load_rows(glpath), load_rows(subpath)
            matched, exceptions = reconcile(gl, sub)
            with connection() as c:
                cur = c.execute('INSERT INTO runs (created_at,gl_file,sub_file,gl_count,sub_count,matched,exceptions,gl_total,sub_total) VALUES (?,?,?,?,?,?,?,?,?)', (datetime.now().isoformat(timespec='seconds'), os.path.basename(glpath), os.path.basename(subpath), len(gl), len(sub), matched, len(exceptions), str(sum(r['amount_cents'] for r in gl)), str(sum(r['amount_cents'] for r in sub))))
                rid = cur.lastrowid
                c.executemany('INSERT INTO transactions (run_id,source,transaction_id,account,txn_date,amount_cents) VALUES (?,?,?,?,?,?)', [(rid, source, r['transaction_id'], r['account'], r['txn_date'], r['amount_cents']) for source, rows in [('GL', gl), ('SUB', sub)] for r in rows])
                c.executemany('INSERT INTO exceptions (run_id,transaction_id,account,issue,gl_count,sub_count,gl_cents,sub_cents) VALUES (?,?,?,?,?,?,?,?)', [(rid, *e) for e in exceptions])
            self.q.put(('success', rid))
        except Exception as e:
            self.q.put(('error', str(e)))

    def poll(self):
        try:
            while True:
                kind, value = self.q.get_nowait()
                if kind == 'error':
                    self.status.set('Error: ' + value)
                    messagebox.showerror('Reconciliation error', value)
                else:
                    self.run_id.set(str(value))
                    self.status.set(f'Reconciliation run {value} completed and saved.')
                    self.refresh_runs()
                    self.load_exceptions()
                    self.load_reconciliation(value)
        except queue.Empty:
            pass
        self.after(100, self.poll)

    def refresh_runs(self):
        self.runs_tree.delete(*self.runs_tree.get_children())
        with connection() as c:
            for r in c.execute('SELECT * FROM runs ORDER BY id DESC LIMIT 200'):
                gl, sub = int(r['gl_total']), int(r['sub_total'])
                self.runs_tree.insert('', 'end', values=(r['id'], r['created_at'], r['gl_count'], r['sub_count'], r['matched'], r['exceptions'], money(gl), money(sub), money(gl-sub)))

    def choose_run(self, _event=None):
        selection = self.runs_tree.selection()
        if selection:
            rid = self.runs_tree.item(selection[0], 'values')[0]
            self.run_id.set(str(rid))
            self.load_exceptions()
            self.load_reconciliation(rid)

    def load_reconciliation(self, rid):
        self.recon_tree.delete(*self.recon_tree.get_children())
        with connection() as c:
            for r in c.execute('SELECT * FROM exceptions WHERE run_id=? ORDER BY id', (rid,)):
                self.recon_tree.insert('', 'end', values=(r['transaction_id'], r['account'], r['issue'], r['gl_count'], r['sub_count'], money(r['gl_cents']), money(r['sub_cents'])))
            run = c.execute('SELECT * FROM runs WHERE id=?', (rid,)).fetchone()
        if run:
            self.result_label.config(text=f"Run {rid}: {run['gl_count']} GL rows | {run['sub_count']} subledger rows | {run['matched']} matched unique keys | {run['exceptions']} exceptions")

    def load_exceptions(self):
        try:
            rid = int(self.run_id.get())
        except ValueError:
            messagebox.showwarning('Select run', 'Enter a numeric run ID or select a dashboard row.')
            return
        self.exc_tree.delete(*self.exc_tree.get_children())
        with connection() as c:
            for r in c.execute('SELECT * FROM exceptions WHERE run_id=? ORDER BY id', (rid,)):
                self.exc_tree.insert('', 'end', values=(r['id'], r['transaction_id'], r['account'], r['issue'], money(r['gl_cents']-r['sub_cents']), r['status'], r['reviewed_by'], r['reviewed_at']))

    def save_review(self):
        selected = self.exc_tree.selection()
        if not selected:
            messagebox.showwarning('Select exception', 'Select an exception row first.')
            return
        reviewer, note = self.reviewer.get().strip(), self.note.get().strip()
        if not reviewer or not note:
            messagebox.showwarning('Review details', 'Reviewer and explanation/evidence reference are required.')
            return
        exc_id = int(self.exc_tree.item(selected[0], 'values')[0])
        with connection() as c:
            c.execute('UPDATE exceptions SET status=?, note=?, reviewed_by=?, reviewed_at=? WHERE id=?', (self.exc_status.get(), note, reviewer, datetime.now().isoformat(timespec='seconds'), exc_id))
        self.load_exceptions()
        self.status.set(f'Exception {exc_id} review saved.')

    def run_sql(self):
        query = self.sql_text.get('1.0', 'end').strip()
        import re
        cleaned = re.sub(r'/\*.*?\*/|--[^\n]*', '', query, flags=re.S).strip()
        if not re.match(r'^(SELECT|WITH)\b', cleaned, flags=re.I) or ';' in cleaned.rstrip(';'):
            messagebox.showerror('Read-only SQL', 'Only one SELECT or WITH query is permitted.')
            return
        try:
            c = connection()
            try:
                c.execute('PRAGMA query_only = ON')
                c.set_progress_handler(lambda: 1, 200000)
                cursor = c.execute(cleaned)
                if cursor.description is None:
                    raise ValueError('Query must return rows')
                columns = [x[0] for x in cursor.description]
                rows = cursor.fetchmany(1001)
                if len(rows) > 1000:
                    rows = rows[:1000]
                    self.status.set('SQL result limited to first 1,000 rows.')
            finally:
                c.close()
            self.sql_result.destroy()
            self.sql_result = self.table(self.frames['SQL Analysis'], columns)
            for row in rows:
                self.sql_result.insert('', 'end', values=[str(x) if x is not None else 'NULL' for x in row])
            if len(rows) <= 1000:
                self.status.set(f'SQL returned {len(rows)} rows.')
        except Exception as e:
            messagebox.showerror('SQL error', str(e))

    def export_report(self):
        if openpyxl is None:
            messagebox.showerror('Missing dependency', 'Install openpyxl: pip install openpyxl')
            return
        try:
            rid = int(self.run_id.get())
        except ValueError:
            messagebox.showwarning('Select run', 'Select a reconciliation run first.')
            return
        with connection() as c:
            run = c.execute('SELECT * FROM runs WHERE id=?', (rid,)).fetchone()
            if not run:
                messagebox.showerror('Run not found', 'No reconciliation run matches this ID.')
                return
            txns = c.execute('SELECT source,transaction_id,account,txn_date,amount_cents FROM transactions WHERE run_id=? ORDER BY id', (rid,)).fetchall()
            exc = c.execute('SELECT transaction_id,account,issue,gl_count,sub_count,gl_cents,sub_cents,status,note,reviewed_by,reviewed_at FROM exceptions WHERE run_id=? ORDER BY id', (rid,)).fetchall()
        path = filedialog.asksaveasfilename(defaultextension='.xlsx', initialfile=f'Reconciliation_Run_{rid}.xlsx', filetypes=[('Excel workbook', '*.xlsx')])
        if not path:
            return
        try:
            from openpyxl.styles import Font, PatternFill
            from openpyxl.utils import get_column_letter
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = 'Run Summary'
            summary = [('Run ID', rid), ('Run Date', run['created_at']), ('GL File', run['gl_file']), ('Subledger File', run['sub_file']), ('GL Rows', run['gl_count']), ('Subledger Rows', run['sub_count']), ('Matched Unique Keys', run['matched']), ('Exception Keys', run['exceptions']), ('GL Total', money(int(run['gl_total']))), ('Subledger Total', money(int(run['sub_total']))), ('GL Less Subledger', money(int(run['gl_total'])-int(run['sub_total']))), ('Control Note', 'Transaction-level comparison; investigate exceptions before financial statement reliance.')]
            for pair in summary:
                ws.append(pair)
            t = wb.create_sheet('Imported Transactions')
            t.append(['Source', 'Transaction ID', 'Account', 'Date', 'Amount'])
            for r in txns:
                t.append([r['source'], r['transaction_id'], r['account'], r['txn_date'], float(Decimal(r['amount_cents'])/100)])
                t.cell(t.max_row, 5).number_format = '#,##0.00;[Red](#,##0.00)'
            e = wb.create_sheet('Exceptions and Reviews')
            e.append(['Transaction ID', 'Account', 'Issue', 'GL Count', 'Sub Count', 'GL Amount', 'Sub Amount', 'Difference', 'Status', 'Explanation', 'Reviewer', 'Reviewed At'])
            for r in exc:
                vals = [r['transaction_id'], r['account'], r['issue'], r['gl_count'], r['sub_count'], float(Decimal(r['gl_cents'])/100), float(Decimal(r['sub_cents'])/100), float(Decimal(r['gl_cents']-r['sub_cents'])/100), r['status'], r['note'], r['reviewed_by'], r['reviewed_at']]
                e.append(vals)
                for col in (6,7,8):
                    e.cell(e.max_row, col).number_format = '#,##0.00;[Red](#,##0.00)'
            for sheet in wb:
                sheet.freeze_panes = 'A2'
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color='FFFFFF')
                    cell.fill = PatternFill('solid', fgColor='234A6B')
                for col in sheet.columns:
                    letter = get_column_letter(col[0].column)
                    sheet.column_dimensions[letter].width = min(55, max(14, max(len(str(cell.value or '')) for cell in col) + 2))
            t.auto_filter.ref = t.dimensions
            e.auto_filter.ref = e.dimensions
            wb.save(path)
            self.status.set(f'Excel report exported: {path}')
            messagebox.showinfo('Export complete', f'Report saved:\n{path}')
        except Exception as err:
            messagebox.showerror('Export failed', str(err))

if __name__ == '__main__':
    Application().mainloop()
